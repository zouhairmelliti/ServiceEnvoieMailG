import os
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def generate_excel(df: pd.DataFrame, output_dir: str = "reports") -> str:
    os.makedirs(output_dir, exist_ok=True)
    filename = f"rapport_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_dir, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Données", index=False)
        ws = writer.sheets["Données"]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

        for col_idx, column in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

            max_len = max(
                df[column].astype(str).map(len).max() if not df.empty else 0,
                len(str(column))
            ) + 2
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 50)

        # Fige la première ligne
        ws.freeze_panes = "A2"

    return filepath